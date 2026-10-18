"""Checks of the skill deck-style-replica's script (run by tests/deck-style-replica-script.test.js when python3 has python-pptx).

A small "company deck" is made here (a logo on every slide, red titles set by hand, old content); then a new deck is built in its style, changed, and checked.
Prints OK, or fails with an AssertionError.
"""
import io
import os
import sys
import tempfile
import zipfile

sys.path.insert(0, sys.argv[1])
import deckstyle as ds
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm

work = tempfile.mkdtemp()
logo, photo = os.path.join(work, "logo.png"), os.path.join(work, "photo.jpg")
Image.new("RGB", (200, 80), (200, 30, 60)).save(logo)
Image.new("RGB", (1600, 900), (40, 120, 180)).save(photo)
prs = Presentation()
for index, title in enumerate(["ACME 2025", "Revenue", "Customers", "Next", "Thanks"]):
    slide = prs.slides.add_slide(prs.slide_layouts[0 if index in (0, 4) else 1])
    slide.shapes.title.text = title
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(0xC8, 0x1E, 0x3C)
    if index not in (0, 4):
        slide.placeholders[1].text_frame.text = "OLD CONTENT 1,234"
    slide.shapes.add_picture(logo, Cm(22), Cm(0.8), Cm(3))
company = os.path.join(work, "company.pptx")
prs.save(company)

# The style is read, the old content is not given unless asked.
info = ds.inspect(company)
assert info["mode_guess"] == "reference", info["mode_guess"]
assert info["recurring_on_slides"] and info["recurring_on_slides"][0]["kind"] == "picture", "the logo is part of the style"
assert info["hand_set_style"]["title"].get("color") == "#C81E3C", "the red titles are part of the style"
assert "OLD CONTENT" not in ds.describe(company)
assert "OLD CONTENT" in ds.describe(company, include_text=True)

# A template file reads as a template.
template = os.path.join(work, "company.potx")
with zipfile.ZipFile(company) as source, zipfile.ZipFile(template, "w") as target:
    for item in source.infolist():
        body = source.read(item.filename)
        if item.filename == "[Content_Types].xml":
            body = body.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
        target.writestr(item, body)
assert ds.inspect(template)["mode_guess"] == "template"

# A new deck: the style comes along, the old content does not, and nothing is left to fix.
out = os.path.join(work, "out", "new.pptx")
report = ds.build(company, {"slides": [
    {"layout": "Title Slide", "title": "Q3 report", "subtitle": "Sales"},
    {"layout": "Title and Content", "title": "Summary", "body": ["Revenue **48.2m**", {"text": "North first", "level": 1}]},
    {"layout": "Two Content", "title": "Pros and cons", "columns": [["Pros", "Stable"], ["Cons", "Costs"]]},
    {"layout": "Title Only", "title": "Monthly", "chart": {"type": "column", "categories": ["Jul", "Aug"], "series": [{"name": "2026", "values": [1, 2]}]}},
    {"layout": "Title Only", "title": "Regions", "table": {"columns": ["Region", "Revenue"], "rows": [["North", "2,100"]]}},
    {"layout": "Picture with Caption", "title": "New store", "image": photo, "body": "Opened in September"},
]}, out)
assert report["problems"] == [], report["problems"]
built = Presentation(out)
assert len(built.slides) == 6
text = " ".join(shape.text_frame.text for slide in built.slides for shape in slide.shapes if shape.has_text_frame)
assert "OLD CONTENT" not in text and "ACME" not in text, "nothing of the old slides"
for slide in built.slides:
    assert any(shape.shape_type == 13 and shape.left > Cm(20) for shape in slide.shapes), "the logo is on every new slide"
    title = slide.shapes.title
    if title is not None and title.text_frame.text:
        assert title.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(0xC8, 0x1E, 0x3C), "titles are red like the old ones"
    for shape in slide.placeholders:
        assert shape.text_frame.text.strip() if shape.has_text_frame else True, "no empty placeholder is left"

# Changes in order, slide files keep unique names, and a layout change puts the text in the main box.
edited = os.path.join(work, "out", "edited.pptx")
result = ds.edit(out, [
    {"op": "move", "slide": 6, "shape": {"placeholder": "title"}, "dx": "1%"},
    {"op": "color", "slide": "all", "role": "title", "color": "accent1"},
    {"op": "duplicate_slide", "slide": 4, "to": 5},
    {"op": "change_layout", "slide": 2, "layout": "Comparison"},
    {"op": "move_slide", "slide": 7, "to": 2},
    {"op": "delete_slide", "slide": 5},
    {"op": "add_slide", "at": 3, "layout": "Section Header", "title": "Part two", "body": "Markets"},
], edited)
assert result["problems"] == [], result["problems"]
with zipfile.ZipFile(edited) as package:
    names = package.namelist()
    assert len(names) == len(set(names)), "no part is written twice"
changed = Presentation(edited)
assert len(changed.slides) == 7
comparison = next(slide for slide in changed.slides if slide.slide_layout.name == "Comparison")
main = next(shape for shape in comparison.placeholders if "Revenue" in shape.text_frame.text)
assert main.left < changed.slide_width / 2 and main.height > changed.slide_height / 3, "the text went to the main box on the left"

# A failed step saves nothing; a missing layout is named with the ones there are.
before = open(edited, "rb").read()
try:
    ds.edit(edited, [{"op": "delete_slide", "slide": 1}, {"op": "move", "slide": 99, "shape": 1}])
    raise AssertionError("a step on a slide that is not there must fail")
except ValueError as error:
    assert "Step 2" in str(error)
assert open(edited, "rb").read() == before
try:
    ds.build(company, {"slides": [{"layout": "Nope", "title": "x"}]}, os.path.join(work, "x.pptx"))
    raise AssertionError("an unknown layout must fail")
except ValueError as error:
    assert "Title and Content" in str(error)

# The check finds what a person would see.
messy = Presentation(out)
box = messy.slides[1].placeholders[1]
box.text_frame.text = "word " * 600
messy.slides[2].shapes.add_textbox(messy.slide_width - Cm(1), Cm(1), Cm(5), Cm(1)).text_frame.text = "off the edge"
messy_path = os.path.join(work, "messy.pptx")
messy.save(messy_path)
found = " ".join(item["problem"] for item in ds.check(messy_path, as_list=True))
assert "may not fit" in found and "past the edge" in found, found
print("OK")
